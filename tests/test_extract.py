import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

import make_samples
from extract import extract_folder, write_xlsx


def test_extraction_and_validation(tmp_path):
    make_samples.main(tmp_path)
    inv = {i["file"]: i for i in extract_folder(tmp_path)}
    assert len(inv) == 5

    a = inv["acme_supplies.pdf"]
    assert a["vendor"] == "Acme Office Supplies Ltd"
    assert a["invoice_number"] == "INV-2026-0101"
    assert (a["date"], a["due_date"], a["currency"]) == ("2026-03-04", "2026-04-03", "GBP")
    assert (a["subtotal"], a["tax"], a["total"]) == (346.6, 69.32, 415.92)
    assert len(a["line_items"]) == 3 and a["line_items"][0]["qty"] == 10
    assert a["status"] == "OK"

    bad = inv["mismatch_total.pdf"]
    assert bad["status"] == "NEEDS REVIEW" and "!= total" in bad["reason"]
    assert sum(i["status"] == "OK" for i in inv.values()) == 4


def test_xlsx(tmp_path):
    make_samples.main(tmp_path)
    out = tmp_path / "o.xlsx"
    write_xlsx(extract_folder(tmp_path), out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Invoices", "Line items"]
    assert wb["Invoices"].max_row == 6
